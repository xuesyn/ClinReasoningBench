from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import RESULTS_DIR


RESULTS_FILE = RESULTS_DIR / "results_filtered_latest.xlsx"
DEFAULT_MODEL_META = {"id": "unknown", "family": "unknown", "access": "unknown"}

MODEL_META: dict[str, dict[str, str]] = {
    "GPT-5.2": {"id": "gpt-5-2", "family": "frontier", "access": "closed"},
    "Claude 4.5 sonnet": {"id": "claude-4-5-sonnet", "family": "frontier", "access": "closed"},
    "Gemini 3 pro preview": {"id": "gemini-3-pro-preview", "family": "frontier", "access": "closed"},
    "Baichuan-M3 (plus)": {"id": "baichuan-m3-plus", "family": "medical", "access": "open"},
    "medgemma": {"id": "medgemma", "family": "medical", "access": "open"},
    "GLM-4.7": {"id": "glm-4-7", "family": "general", "access": "open"},
    "kimi-k2.5": {"id": "kimi-k2-5", "family": "general", "access": "open"},
    "Deepseek-v3.2": {"id": "deepseek-v3-2", "family": "general", "access": "open"},
    "Qwen-Max": {"id": "qwen-max", "family": "general", "access": "closed"},
    "Qwen3.5-Plus": {"id": "qwen-3-5-plus", "family": "general", "access": "open"},
}

DATASET_META: dict[str, dict[str, str]] = {
    "cnlc": {"label": "CNLC", "task_family": "oncology"},
    "bclc": {"label": "BCLC", "task_family": "oncology"},
    "nsclc": {"label": "NSCLC", "task_family": "oncology"},
    "acs": {"label": "Acute Coronary Syndromes", "task_family": "acute"},
    "ais": {"label": "Acute Ischemic Stroke", "task_family": "acute"},
}

_CACHE: dict[str, Any] = {"mtime_ns": None, "payload": None}


def get_results_dashboard() -> dict[str, Any]:
    if not RESULTS_FILE.exists():
        raise KeyError(f"Results file not found: {RESULTS_FILE}")

    mtime_ns = RESULTS_FILE.stat().st_mtime_ns
    if _CACHE["mtime_ns"] == mtime_ns and _CACHE["payload"] is not None:
        return _CACHE["payload"]

    payload = _build_results_dashboard(RESULTS_FILE)
    _CACHE["mtime_ns"] = mtime_ns
    _CACHE["payload"] = payload
    return payload


def _build_results_dashboard(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        raise KeyError(f"Results file is empty: {path}")

    header = rows[0]
    models = [str(value).strip() for value in header[4:] if str(value or "").strip()]
    dataset_blocks = _parse_dataset_blocks(rows[1:], models)
    metric_catalog = _build_metric_catalog(dataset_blocks)

    return {
        "source_file": str(path),
        "models": [_build_model_entry(model) for model in models],
        "datasets": dataset_blocks,
        "metric_catalog": metric_catalog,
    }


def _build_model_entry(model: str) -> dict[str, str]:
    meta = MODEL_META.get(model)
    if meta is not None:
        return {"name": model, **meta}
    return {"name": model, **DEFAULT_MODEL_META, "id": _slugify(model)}


def _parse_dataset_blocks(rows: list[tuple[Any, ...]], models: list[str]) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    last_dataset = ""
    last_disease = ""
    last_guideline = ""

    for row in rows:
        padded = list(row) + [None] * max(0, 4 + len(models) - len(row))
        dataset_raw = _clean_text(padded[0])
        disease_raw = _clean_text(padded[1])
        guideline_raw = _clean_text(padded[2])
        metric_raw = _clean_text(padded[3])

        if dataset_raw:
            last_dataset = dataset_raw
        if disease_raw:
            last_disease = disease_raw
        if guideline_raw and not _looks_like_sample_size(guideline_raw):
            last_guideline = guideline_raw

        if metric_raw == "总榜":
            dataset_id = _resolve_dataset_id(last_dataset, last_disease, last_guideline)
            meta = DATASET_META[dataset_id]
            current = {
                "id": dataset_id,
                "label": meta["label"],
                "source": last_dataset or "-",
                "disease": last_disease or "-",
                "guideline": last_guideline or "-",
                "task_family": meta["task_family"],
                "sample_size": None,
                "metrics": [],
            }
            blocks.append(current)
            continue

        if current is None:
            continue

        sample_candidate = guideline_raw or disease_raw or dataset_raw
        if not metric_raw and _looks_like_sample_size(sample_candidate):
            current["sample_size"] = _parse_sample_size(sample_candidate)
            continue

        if not metric_raw:
            continue

        metric_def = _classify_metric(metric_raw)
        values = {
            model: float(value)
            for model, value in zip(models, padded[4 : 4 + len(models)])
            if isinstance(value, (int, float))
        }
        if not values:
            continue

        current["metrics"].append(
            {
                "key": metric_def["key"],
                "label": metric_def["label"],
                "raw_label": metric_raw,
                "group": metric_def["group"],
                "direction": metric_def["direction"],
                "values": values,
            }
        )

    return blocks


def _build_metric_catalog(dataset_blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    catalog: dict[str, dict[str, Any]] = {}
    for block in dataset_blocks:
        for metric in block["metrics"]:
            if metric["key"] not in catalog:
                catalog[metric["key"]] = {
                    "key": metric["key"],
                    "label": metric["label"],
                    "group": metric["group"],
                    "direction": metric["direction"],
                    "datasets": [],
                }
            catalog[metric["key"]]["datasets"].append(block["id"])

    return sorted(catalog.values(), key=lambda item: (item["group"], item["label"]))


def _resolve_dataset_id(dataset: str, disease: str, guideline: str) -> str:
    dataset_low = dataset.lower()
    disease_low = disease.lower()
    guideline_low = guideline.lower()

    if "cnlc" in guideline_low:
        return "cnlc"
    if "bclc" in guideline_low:
        return "bclc"
    if "肺癌" in disease or "nsclc" in guideline_low:
        return "nsclc"
    if "心梗" in disease or "cir.0000000000001309" in guideline_low:
        return "acs"
    if "脑梗" in disease or "strokeaha.119.027708" in guideline_low:
        return "ais"
    if dataset_low == "seer" and "肝癌" in disease:
        return "cnlc"
    if dataset_low == "mimic" and "心" in disease:
        return "acs"
    if dataset_low == "mimic" and "脑" in disease:
        return "ais"
    raise KeyError(f"Unknown dataset block: dataset={dataset!r}, disease={disease!r}, guideline={guideline!r}")


def _classify_metric(raw_label: str) -> dict[str, str]:
    label = raw_label.strip()
    lowered = label.lower()

    if lowered == "parse_error_rate":
        return {"key": "parse_error_rate", "label": "Parse Error Rate", "group": "robustness", "direction": "lower"}
    if label == "PW ECE Mem":
        return {"key": "pw_ece_mem", "label": "PW-ECE Mem", "group": "calibration", "direction": "lower"}
    if label == "PW ECE RT":
        return {"key": "pw_ece_rt", "label": "PW-ECE RT", "group": "calibration", "direction": "lower"}
    if label == "NDCG ECE":
        return {"key": "ndcg_ece", "label": "NDCG ECE", "group": "calibration", "direction": "lower"}
    if label == "IOU ECE":
        return {"key": "iou_ece", "label": "IoU ECE", "group": "calibration", "direction": "lower"}
    if label == "F1 ECE":
        return {"key": "f1_ece", "label": "F1 ECE", "group": "calibration", "direction": "lower"}
    if "完全匹配" in label:
        return {"key": "indication_exact", "label": "Indication/Contra Exact Match", "group": "labeling", "direction": "higher"}
    if "禁忌证和适应证" in label:
        return {"key": "indication_hit", "label": "Indication/Contra Hit Score", "group": "labeling", "direction": "higher"}
    if label == "治疗 ndcg":
        return {"key": "treatment_ndcg", "label": "Treatment NDCG", "group": "treatment", "direction": "higher"}
    if "治疗（proc_acc）" in label:
        return {"key": "procedure_acc", "label": "Procedure Accuracy", "group": "treatment", "direction": "higher"}
    if label == "proc_recall":
        return {"key": "procedure_recall", "label": "Procedure Recall", "group": "treatment", "direction": "higher"}
    if label == "med_acc":
        return {"key": "med_acc", "label": "Medication Accuracy", "group": "treatment", "direction": "higher"}
    if label == "med_recall":
        return {"key": "med_recall", "label": "Medication Recall", "group": "treatment", "direction": "higher"}
    if label == "all_acc":
        return {"key": "all_acc", "label": "Overall Treatment Accuracy", "group": "treatment", "direction": "higher"}
    if label == "all_recall":
        return {"key": "all_recall", "label": "Overall Treatment Recall", "group": "treatment", "direction": "higher"}
    if "graph+citation" in lowered or "推理" in label:
        return {"key": "reasoning_score", "label": "Reasoning (Graph + Citation)", "group": "reasoning", "direction": "higher"}
    if "[tnm]" in lowered:
        return {"key": "staging_tnm", "label": "Staging Accuracy (TNM)", "group": "staging", "direction": "higher"}
    if "分期" in label:
        return {"key": "staging_accuracy", "label": "Staging Accuracy", "group": "staging", "direction": "higher"}

    return {"key": _slugify(label), "label": label, "group": "other", "direction": "higher"}


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float)):
        if int(value) == value:
            return str(int(value))
        return str(value)
    return str(value).strip()


def _looks_like_sample_size(value: str) -> bool:
    stripped = value.strip()
    if not stripped:
        return False
    return stripped.startswith("-") and stripped[1:].isdigit()


def _parse_sample_size(value: str) -> int | None:
    if not _looks_like_sample_size(value):
        return None
    return abs(int(value))


def _slugify(value: str) -> str:
    chars = []
    for char in value.lower():
        if char.isalnum():
            chars.append(char)
        else:
            chars.append("-")
    return "-".join(filter(None, "".join(chars).split("-")))
